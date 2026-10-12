"""episodeのJSONから、確認用の台本（Word）を書き出す。

使い方:
  python make_script_docx.py episode01.json 台本.docx [--timeline timeline.txt]
"""
import argparse
import json
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAMES = {"K": "工藤", "N": "中山", "NA": "中野", "S": "佐伯（声）", "NR": "ナレーション"}
SCENE = {"talk": "会話", "board": "図解", "list": "AI案一覧", "mail": "メール画面", "card": "タイトル",
         "chapter": "章扉", "kv": "キービジュアル", "summary": "まとめ", "cta": "相談導線", "next": "次回予告"}


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def screen_text(sh):
    t = sh["type"]
    if t == "board":
        rows = [f"{a}　{b}" for a, b in sh["items"]]
        return f"【{sh['heading']}】\n" + "\n".join(rows)
    if t == "summary":
        return f"【{sh['heading']}】\n" + "\n".join(f"{i + 1}. {x}" for i, x in enumerate(sh["items"]))
    if t == "mail":
        return f"差出人　{sh['from']}\n件名　{sh['subject']}"
    if t == "card":
        return f"{sh['sub']}　{sh['big']}"
    if t == "chapter":
        return f"{sh['no']}　{sh['title']}"
    if t in ("cta", "next"):
        return sh["big"].replace("\n", "") + (f"\n{sh['sub']}" if sh.get("sub") else "")
    if t == "list":
        return "AIの施策案32件を一覧表示"
    return sh.get("telop_note", "")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("episode")
    ap.add_argument("out")
    ap.add_argument("--timeline")
    a = ap.parse_args()
    ep = json.load(open(a.episode, encoding="utf-8"))
    times = {}
    if a.timeline:
        for ln in open(a.timeline, encoding="utf-8"):
            m = re.match(r"(\d+:\d+)\.\d+\s+#(\d+)", ln)
            if m:
                times[int(m.group(2))] = m.group(1)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin = sec.bottom_margin = Mm(14)
    st = doc.styles["Normal"]
    st.font.name = "Yu Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    st.font.size = Pt(10)

    h = doc.add_paragraph()
    r = h.add_run(ep["title"])
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph("台本（Vコンテ対応版）　" + ep["note"]).runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Mm(12), Mm(16), Mm(24), Mm(26), Mm(189)]
    tbl.autofit = False
    for col, w in zip(tbl.columns, widths):
        col.width = w
    for c, (txt, w) in enumerate(zip(["No.", "時刻", "場面", "話者", "台詞／画面"], widths)):
        cell = tbl.rows[0].cells[c]
        cell.width = w
        cell.text = txt
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "1D1F24")
    for i, sh in enumerate(ep["shots"]):
        row = tbl.add_row().cells
        for c, w in enumerate(widths):
            row[c].width = w
        row[0].text = f"{i + 1:02d}"
        row[1].text = times.get(i, "")
        row[2].text = "カットイン" if sh.get("cutin") else SCENE[sh["type"]]
        row[3].text = NAMES.get(sh.get("sp"), "")
        p = row[4].paragraphs[0]
        line = "" if sh.get("silent") else sh.get("text", "")
        if line and sh["type"] not in ("mail",):
            p.add_run(line).bold = sh["type"] in ("talk", "board", "list")
        elif sh["type"] == "mail":
            p.add_run(line)
        scr = screen_text(sh)
        if scr:
            rr = (row[4].add_paragraph() if p.runs else p).add_run("画面：" + scr)
            rr.font.size = Pt(9)
            rr.font.color.rgb = RGBColor(0x8A, 0x32, 0x2A)
        if sh["type"] == "chapter":
            for c in row:
                shade(c, "F3ECDC")
    doc.add_paragraph()
    doc.add_paragraph("仮音声の配役").runs[0].bold = True
    for t in ["中山　JVNV-M1を低め・ゆっくり・抑揚強めに調整",
              "ナレーション　JVNV-M1をさらに低く、抑揚を抑えて読み上げ調に",
              "工藤　JVNV-M2を少し低め、抑揚控えめで淡々と",
              "中野　JVNV-M2を少し高め・速め、明るい感情スタイルで",
              "佐伯（架空の依頼主）　JVNV-F1",
              "音声モデルはStyle-Bert-VITS2のJVNVモデル（CC BY-SA 4.0）。公開する場合はクレジット表記が必要。本番は収録か正式な音声に差し替える。"]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph()
    doc.add_paragraph("確認事項").runs[0].bold = True
    for t in ["中山さんの語尾と言い回しは仮。本人の話し方に合わせて調整する。",
              "中野さんの立ち絵とキャラ設定が届いたら、シルエットを差し替える。",
              "ひなた食品、佐伯さん、商品名「まいにちポタ」はすべて架空。実案件に寄せる場合は守秘範囲を確認する。",
              "相談導線の診断ページは未制作。公開前にURLと文言を確定する。",
              "中山さんの関西弁は、合成音声だと標準語のアクセントで読まれる。収録時に本人の抑揚で直す。",
              "BGMは仮の自動生成。公開前に権利処理済みの曲へ差し替える。"]:
        doc.add_paragraph(t, style="List Bullet")
    doc.save(a.out)


if __name__ == "__main__":
    main()
